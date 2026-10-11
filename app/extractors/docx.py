"""Word (.docx) extractor (uses python-docx; degrades gracefully)."""
from __future__ import annotations

import io

from app.extractors.base import Extractor, ExtractResult


class DocxExtractor(Extractor):
    source_type = "docx"

    def extract(self, content) -> ExtractResult:  # noqa: ANN001
        if not isinstance(content, (bytes, bytearray)):
            raise ValueError("DocxExtractor expects bytes")

        try:
            import docx  # python-docx
        except ModuleNotFoundError:
            return ExtractResult(
                text="",
                source_type=self.source_type,
                meta={"error": "python-docx not installed; run pip install python-docx"},
            )

        document = docx.Document(io.BytesIO(content))
        parts: list[str] = [p.text for p in document.paragraphs if p.text]

        # Table cells are a common hiding spot for injected instructions.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())

        return ExtractResult(
            text="\n".join(parts).strip(),
            source_type=self.source_type,
        )
